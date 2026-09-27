#!/usr/bin/env python
"""Compare the two wrist-view presentation conditions at the participant level.

Episodes are first averaged within participant/seed and then within participant, so repeat
trials are never treated as independent samples. The confirmed within-participant design is
paired by default. Independent-samples tests are available only through the explicit
``--independent_diagnostic`` switch and must not be reported as the formal study result.
Binary episode outcomes follow the same paired/diagnostic design choice.

For protocol v2, the script fits an effect-coded condition x task factorial OLS model with
``condition_order`` adjustment. Episode repeats are averaged to participant-condition-task
cells and participant fixed effects are included. This is a transparent repeated-measures
approximation using NumPy/SciPy, not a full mixed model.
"""

from __future__ import annotations

import argparse
from pathlib import Path

import numpy as np
import pandas as pd
from scipy import stats

try:
    from .analysis_reproducibility import (
        fingerprint_inputs,
        fingerprint_path,
        git_state,
        runtime_environment,
        write_json,
    )
except ImportError:  # Direct script execution.
    from analysis_reproducibility import (  # type: ignore[no-redef]
        fingerprint_inputs,
        fingerprint_path,
        git_state,
        runtime_environment,
        write_json,
    )

NON_FEATURE = {
    "episode",
    "n_frames",
    "success",
    "failure_type",
    "condition",
    "display",
    "split",
    "pair_id",
    "participant",
    "participant_id",
    "task",
    "task_id",
    "trial",
    "trial_index",
    "repetition",
    "session",
    "condition_order",
    "quality_status",
    "feature_usable",
    "completion_time_source",
    "reported_completion_time_s",
    "raw_duration_s",
    "seed",
}
MISSING_IDENTIFIER_TOKENS = {"", "nan", "none", "<na>", "null"}


def cohens_d(a: np.ndarray, b: np.ndarray) -> float:
    """Independent-samples standardized mean difference."""
    if len(a) < 2 or len(b) < 2:
        return float("nan")
    pooled_variance = ((len(a) - 1) * a.var(ddof=1) + (len(b) - 1) * b.var(ddof=1)) / (len(a) + len(b) - 2)
    pooled_sd = np.sqrt(pooled_variance)
    return float((a.mean() - b.mean()) / pooled_sd) if pooled_sd > 0 else float("nan")


def cohens_dz(a: np.ndarray, b: np.ndarray) -> float:
    """Paired-samples standardized mean difference (mean difference / SD difference)."""
    difference = a - b
    if len(difference) < 2:
        return float("nan")
    sd = difference.std(ddof=1)
    if sd == 0:
        return float("nan")
    return float(difference.mean() / sd)


def effect_label(effect_size: float) -> str:
    if np.isnan(effect_size):
        return "n/a"
    absolute = abs(effect_size)
    if absolute < 0.2:
        return "negligible"
    if absolute < 0.5:
        return "small"
    if absolute < 0.8:
        return "medium"
    return "large"


def _add_fdr(table: pd.DataFrame) -> pd.DataFrame:
    if table.empty:
        return table
    adjusted = np.full(len(table), np.nan)
    finite = np.isfinite(table["p"].to_numpy(dtype=float))
    if finite.any():
        adjusted[finite] = stats.false_discovery_control(
            table.loc[finite, "p"].to_numpy(dtype=float), method="bh"
        )
    table = table.copy()
    table["p_fdr"] = adjusted
    table["significant"] = table["p_fdr"] < 0.05
    return table.sort_values("p_fdr", na_position="last").reset_index(drop=True)


def _validate_pair_keys(
    df_a: pd.DataFrame,
    df_b: pd.DataFrame,
    pair_keys: list[str],
) -> None:
    if not pair_keys:
        raise ValueError("paired analysis requires at least one pair key")
    missing = [key for key in pair_keys if key not in df_a.columns or key not in df_b.columns]
    if missing:
        raise ValueError(f"pair keys are absent from one or both feature tables: {missing}")
    for name, dataframe in (("A", df_a), ("B", df_b)):
        if dataframe[pair_keys].isna().any(axis=None):
            raise ValueError(f"Condition {name} has missing pair-key values")
        duplicated = dataframe.duplicated(pair_keys, keep=False)
        if duplicated.any():
            examples = dataframe.loc[duplicated, pair_keys].drop_duplicates().head(5).to_dict("records")
            raise ValueError(
                f"Condition {name} pair keys are not unique; add trial/repetition keys. Examples: {examples}"
            )
    keys_a = set(df_a[pair_keys].itertuples(index=False, name=None))
    keys_b = set(df_b[pair_keys].itertuples(index=False, name=None))
    if keys_a != keys_b:
        only_a = sorted(keys_a - keys_b, key=str)[:10]
        only_b = sorted(keys_b - keys_a, key=str)[:10]
        raise ValueError(
            "paired tables do not contain exactly the same matched trials; "
            f"only_A={only_a}, only_B={only_b}. Refusing an inner join that would drop pairs."
        )


def _paired_values(
    df_a: pd.DataFrame,
    df_b: pd.DataFrame,
    feature: str,
    pair_keys: list[str],
) -> tuple[np.ndarray, np.ndarray]:
    merged = df_a[pair_keys + [feature]].merge(
        df_b[pair_keys + [feature]],
        on=pair_keys,
        how="inner",
        suffixes=("_A", "_B"),
        validate="one_to_one",
    )
    merged = merged.dropna(subset=[f"{feature}_A", f"{feature}_B"])
    return (
        merged[f"{feature}_A"].to_numpy(dtype=float),
        merged[f"{feature}_B"].to_numpy(dtype=float),
    )


def _paired_ttest(a: np.ndarray, b: np.ndarray) -> tuple[float, float]:
    difference = a - b
    if np.allclose(difference, 0):
        return 0.0, 1.0
    if difference.std(ddof=1) == 0:
        return float(np.copysign(np.inf, difference.mean())), 0.0
    result = stats.ttest_rel(a, b)
    return float(result.statistic), float(result.pvalue)


def _mean_ci95(values: np.ndarray) -> tuple[float, float]:
    """Return a descriptive two-sided 95% t interval for one mean."""
    values = np.asarray(values, dtype=float)
    if len(values) < 2:
        return float("nan"), float("nan")
    mean = float(values.mean())
    standard_error = float(values.std(ddof=1) / np.sqrt(len(values)))
    if np.isclose(standard_error, 0.0):
        return mean, mean
    margin = float(stats.t.ppf(0.975, len(values) - 1) * standard_error)
    return mean - margin, mean + margin


def _welch_mean_difference_ci95(a: np.ndarray, b: np.ndarray) -> tuple[float, float]:
    """Return a two-sided Welch 95% interval for mean(A)-mean(B)."""
    difference = float(a.mean() - b.mean())
    variance_a = float(a.var(ddof=1) / len(a))
    variance_b = float(b.var(ddof=1) / len(b))
    standard_error_squared = variance_a + variance_b
    if np.isclose(standard_error_squared, 0.0):
        return difference, difference
    denominator = variance_a**2 / (len(a) - 1) + variance_b**2 / (len(b) - 1)
    degrees_of_freedom = standard_error_squared**2 / denominator
    margin = float(stats.t.ppf(0.975, degrees_of_freedom) * np.sqrt(standard_error_squared))
    return difference - margin, difference + margin


def validate_condition_table(
    dataframe: pd.DataFrame,
    expected_condition: str,
) -> None:
    """Reject mislabeled/swapped inputs instead of overwriting their condition."""
    required = {"condition", "participant_id", "success", "condition_order", "task_id"}
    missing = sorted(required - set(dataframe.columns))
    if missing:
        raise ValueError(f"{expected_condition} table is missing columns: {missing}")
    if dataframe["condition"].isna().any():
        raise ValueError("condition must be populated for every row")
    condition_values = dataframe["condition"].astype(str).str.strip()
    if condition_values.str.lower().isin(MISSING_IDENTIFIER_TOKENS).any():
        raise ValueError("condition must be populated for every row")
    conditions = set(condition_values)
    if conditions != {expected_condition}:
        raise ValueError(
            f"expected every row to have condition={expected_condition!r}; found={sorted(conditions)}"
        )
    if dataframe["participant_id"].isna().any():
        raise ValueError("participant_id must be populated for every row")
    participants = dataframe["participant_id"].astype(str).str.strip()
    if participants.str.lower().isin(MISSING_IDENTIFIER_TOKENS).any():
        raise ValueError("participant_id must be populated for every row")
    for field in ("task_id", "seed", "split"):
        if field not in dataframe.columns:
            continue
        if dataframe[field].isna().any():
            raise ValueError(f"{field} must be populated for every row when present")
        values = dataframe[field].astype(str).str.strip()
        if values.str.lower().isin(MISSING_IDENTIFIER_TOKENS).any():
            raise ValueError(f"{field} must be populated for every row when present")


def aggregate_participant_features(
    dataframe: pd.DataFrame,
    features: list[str],
) -> pd.DataFrame:
    """Give tasks equal weight, then average seeds within each participant."""
    group_columns = ["participant_id"]
    if "seed" in dataframe.columns:
        seed_values = dataframe["seed"]
        if (
            seed_values.isna().any()
            or seed_values.astype(str).str.strip().str.lower().isin(MISSING_IDENTIFIER_TOKENS).any()
        ):
            raise ValueError("seed must be populated for all rows when present")
        group_columns.append("seed")
    if "task_id" in dataframe.columns:
        task_levels = set(dataframe["task_id"].astype(str).unique())
        task_counts = dataframe.groupby(group_columns, observed=True)["task_id"].nunique()
        if not task_counts.eq(len(task_levels)).all():
            raise ValueError("every participant/seed must contain every task before equal-task aggregation")
        task_level = dataframe.groupby(
            [*group_columns, "task_id"], as_index=False, observed=True
        )[features].mean()
        seed_level = task_level.groupby(group_columns, as_index=False, observed=True)[features].mean()
    else:
        seed_level = dataframe.groupby(group_columns, as_index=False, observed=True)[features].mean()
    participant_level = seed_level.groupby("participant_id", as_index=False, observed=True)[features].mean()
    if len(participant_level) < 2:
        raise ValueError(
            "formal continuous inference requires at least 2 participants; "
            f"found {participant_level['participant_id'].tolist()}"
        )
    return participant_level


def continuous_comparison(
    df_a: pd.DataFrame,
    df_b: pd.DataFrame,
    features: list[str],
    *,
    paired: bool = False,
    pair_keys: list[str] | None = None,
) -> pd.DataFrame:
    """Compare continuous features with Welch or matched-pair t-tests."""
    if paired:
        _validate_pair_keys(df_a, df_b, pair_keys or [])

    rows: list[dict[str, object]] = []
    for feature in features:
        if paired:
            a, b = _paired_values(df_a, df_b, feature, pair_keys or [])
            if len(a) < 2:
                raise ValueError(
                    f"feature {feature!r} has only {len(a)} complete participant pair(s); "
                    "at least 2 are required"
                )
            statistic, p_value = _paired_ttest(a, b)
            effect_size = cohens_dz(a, b)
            test_name = "paired_t"
            n_pairs: int | float = len(a)
            ci_low, ci_high = _mean_ci95(a - b)
        else:
            a = df_a[feature].dropna().to_numpy(dtype=float)
            b = df_b[feature].dropna().to_numpy(dtype=float)
            if len(a) < 2 or len(b) < 2:
                raise ValueError(
                    f"feature {feature!r} has fewer than 2 complete participants (A={len(a)}, B={len(b)})"
                )
            result = stats.ttest_ind(a, b, equal_var=False)
            statistic, p_value = float(result.statistic), float(result.pvalue)
            effect_size = cohens_d(a, b)
            test_name = "welch_t"
            n_pairs = np.nan
            ci_low, ci_high = _welch_mean_difference_ci95(a, b)

        rows.append(
            {
                "feature": feature,
                "test": test_name,
                "n_A": len(a),
                "n_B": len(b),
                "n_pairs": n_pairs,
                "A_mean": a.mean(),
                "A_std": a.std(ddof=1),
                "B_mean": b.mean(),
                "B_std": b.std(ddof=1),
                "diff_A_minus_B": a.mean() - b.mean(),
                "diff_ci95_low": ci_low,
                "diff_ci95_high": ci_high,
                "t": statistic,
                "p": p_value,
                "effect_size": effect_size,
                "effect_size_type": "cohens_dz" if paired else "cohens_d",
                # Retain the old column for downstream spreadsheets.
                "cohens_d": effect_size,
                "effect": effect_label(effect_size),
            }
        )
    return _add_fdr(pd.DataFrame(rows))


def _binary_success(series: pd.Series, condition: str) -> pd.Series:
    if series.isna().any():
        raise ValueError(f"Condition {condition} success has {int(series.isna().sum())} missing value(s)")
    values = pd.to_numeric(series, errors="coerce")
    nonnumeric = values.isna()
    if nonnumeric.any():
        bad = series.loc[nonnumeric].astype(str).unique().tolist()
        raise ValueError(f"Condition {condition} success contains non-numeric values: {bad}")
    invalid = ~values.isin([0, 1])
    if invalid.any():
        raise ValueError(
            f"Condition {condition} success must contain only 0/1; "
            f"invalid={sorted(values.loc[invalid].unique().tolist())}"
        )
    return values.astype(int)


def participant_success_comparison(
    df_a: pd.DataFrame,
    df_b: pd.DataFrame,
    *,
    paired: bool = True,
) -> dict[str, object]:
    """Compare participant success proportions under the selected design."""
    for condition, dataframe in (
        ("A_mobile_colocated", df_a),
        ("B_desktop_separated", df_b),
    ):
        _binary_success(dataframe["success"], condition)

    def participant_rates(dataframe: pd.DataFrame) -> pd.DataFrame:
        columns = ["participant_id"]
        if "seed" in dataframe.columns:
            seed_values = dataframe["seed"]
            if (
                seed_values.isna().any()
                or seed_values.astype(str).str.strip().str.lower().isin(MISSING_IDENTIFIER_TOKENS).any()
            ):
                raise ValueError("seed must be populated for all success rows")
            columns.append("seed")
        normalized = dataframe.assign(success=pd.to_numeric(dataframe["success"]))
        if "task_id" in normalized.columns:
            task_levels = set(normalized["task_id"].astype(str).unique())
            task_counts = normalized.groupby(columns, observed=True)["task_id"].nunique()
            if not task_counts.eq(len(task_levels)).all():
                raise ValueError("every participant/seed must contain every task for success aggregation")
            task_rates = normalized.groupby(
                [*columns, "task_id"], as_index=False, observed=True
            )["success"].mean()
            seed_rates = task_rates.groupby(columns, as_index=False, observed=True)["success"].mean()
        else:
            seed_rates = normalized.groupby(columns, as_index=False, observed=True)["success"].mean()
        return seed_rates.groupby("participant_id", as_index=False, observed=True)["success"].mean()

    rates_a, rates_b = participant_rates(df_a), participant_rates(df_b)
    if not paired:
        if len(rates_a) < 2 or len(rates_b) < 2:
            raise ValueError("independent success diagnostic requires at least 2 participants per condition")
        result = stats.ttest_ind(rates_a["success"], rates_b["success"], equal_var=False)
        ci_low, ci_high = _welch_mean_difference_ci95(
            rates_a["success"].to_numpy(dtype=float),
            rates_b["success"].to_numpy(dtype=float),
        )
        return {
            "test": "participant_welch_t_diagnostic",
            "SR_A": float(rates_a["success"].mean()),
            "SR_B": float(rates_b["success"].mean()),
            "n_a": len(rates_a),
            "n_b": len(rates_b),
            "n_pairs": np.nan,
            "SR_A_minus_B": float(rates_a["success"].mean() - rates_b["success"].mean()),
            "diff_ci95_low": ci_low,
            "diff_ci95_high": ci_high,
            "p": float(result.pvalue),
            "limitation": (
                "Diagnostic independent-samples comparison of participant success "
                "proportions; it does not represent the preregistered paired design."
            ),
        }

    participants_a = set(rates_a["participant_id"])
    participants_b = set(rates_b["participant_id"])
    if participants_a != participants_b:
        raise ValueError(
            "participant success sets differ; "
            f"only_A={sorted(participants_a - participants_b)}, "
            f"only_B={sorted(participants_b - participants_a)}"
        )
    if len(participants_a) < 2:
        raise ValueError("paired success inference requires at least 2 complete participants")
    matched = rates_a.merge(
        rates_b,
        on="participant_id",
        how="inner",
        suffixes=("_A", "_B"),
        validate="one_to_one",
    )
    differences = matched["success_A"] - matched["success_B"]
    ci_low, ci_high = _mean_ci95(differences.to_numpy(dtype=float))
    nonzero = differences[~np.isclose(differences, 0)]
    positive = int((nonzero > 0).sum())
    p_value = float(stats.binomtest(positive, n=len(nonzero), p=0.5).pvalue) if len(nonzero) else 1.0
    return {
        "test": "participant_exact_sign",
        "SR_A": float(matched["success_A"].mean()),
        "SR_B": float(matched["success_B"].mean()),
        "n_a": len(matched),
        "n_b": len(matched),
        "n_pairs": len(matched),
        "SR_A_minus_B": float(differences.mean()),
        "diff_ci95_low": ci_low,
        "diff_ci95_high": ci_high,
        "n_discordant_participants": len(nonzero),
        "p": p_value,
        "limitation": (
            "Exact sign test on participant success proportions; a binomial mixed model "
            "requires an additional statistical-model dependency."
        ),
    }


def _effect_codes(values: pd.Series) -> tuple[np.ndarray, list[object]]:
    """Return sum-to-zero codes and deterministically ordered levels."""
    levels = sorted(values.dropna().unique().tolist(), key=str)
    if len(levels) < 2:
        return np.empty((len(values), 0)), levels
    codes = np.zeros((len(values), len(levels) - 1), dtype=float)
    last = levels[-1]
    for column, level in enumerate(levels[:-1]):
        codes[:, column] = (values == level).astype(float) - (values == last).astype(float)
    return codes, levels


def _factorial_design(
    frame: pd.DataFrame,
    subject_key: str | None,
    factor_column: str = "task_id",
) -> tuple[np.ndarray, dict[str, list[int]], dict[str, list[object]]]:
    condition_codes, condition_levels = _effect_codes(frame["condition"])
    factor_codes, factor_levels = _effect_codes(frame[factor_column])
    parts = [np.ones((len(frame), 1), dtype=float)]
    offset = 1

    if subject_key is not None:
        subject_codes, _ = _effect_codes(frame[subject_key])
        if subject_codes.shape[1]:
            parts.append(subject_codes)
            offset += subject_codes.shape[1]

    terms: dict[str, list[int]] = {}
    order = pd.to_numeric(frame["condition_order"], errors="coerce")
    if order.isna().any() or not np.isfinite(order).all():
        raise ValueError("condition_order must be finite numeric values")
    order_codes = (order.to_numpy(dtype=float) - float(order.mean())).reshape(-1, 1)
    parts.append(order_codes)
    terms["condition_order"] = [offset]
    offset += 1

    parts.append(condition_codes)
    terms["condition"] = list(range(offset, offset + condition_codes.shape[1]))
    offset += condition_codes.shape[1]
    parts.append(factor_codes)
    terms[factor_column] = list(range(offset, offset + factor_codes.shape[1]))
    offset += factor_codes.shape[1]

    interaction = (condition_codes[:, :, None] * factor_codes[:, None, :]).reshape(len(frame), -1)
    parts.append(interaction)
    terms[f"condition:{factor_column}"] = list(range(offset, offset + interaction.shape[1]))
    return (
        np.concatenate(parts, axis=1),
        terms,
        {"condition": condition_levels, factor_column: factor_levels},
    )


def _partial_f_test(
    design: np.ndarray,
    response: np.ndarray,
    removed_columns: list[int],
) -> tuple[float, float, int, int] | None:
    full_coefficients, _, full_rank, _ = np.linalg.lstsq(design, response, rcond=None)
    full_residual = response - design @ full_coefficients
    full_rss = float(full_residual @ full_residual)
    denominator_df = len(response) - full_rank
    reduced = np.delete(design, removed_columns, axis=1)
    reduced_coefficients, _, reduced_rank, _ = np.linalg.lstsq(reduced, response, rcond=None)
    reduced_residual = response - reduced @ reduced_coefficients
    reduced_rss = float(reduced_residual @ reduced_residual)
    numerator_df = int(full_rank - reduced_rank)
    if numerator_df <= 0 or denominator_df <= 0:
        return None
    numerator = max(0.0, reduced_rss - full_rss) / numerator_df
    denominator = full_rss / denominator_df
    if denominator <= np.finfo(float).eps:
        statistic = float("inf") if numerator > np.finfo(float).eps else 0.0
    else:
        statistic = numerator / denominator
    p_value = float(stats.f.sf(statistic, numerator_df, denominator_df))
    return statistic, p_value, numerator_df, int(denominator_df)


def factorial_condition_factor(
    dataframe: pd.DataFrame,
    features: list[str],
    *,
    subject_key: str | None = None,
    factor_column: str,
) -> pd.DataFrame:
    """Fit a condition x task model, optionally blocking by participant."""
    if "condition" not in dataframe.columns or factor_column not in dataframe.columns:
        return pd.DataFrame()
    if "condition_order" not in dataframe.columns:
        raise ValueError(f"condition_order is required to adjust condition x {factor_column}")
    if dataframe["condition"].nunique() < 2 or dataframe[factor_column].nunique() < 2:
        return pd.DataFrame()
    if subject_key is not None and subject_key not in dataframe.columns:
        raise ValueError(f"subject key {subject_key!r} is absent from the feature tables")

    rows: list[dict[str, object]] = []
    for feature in features:
        columns = ["condition", factor_column, "condition_order", feature]
        if subject_key is not None:
            columns.append(subject_key)
        has_seed = "seed" in dataframe.columns
        if has_seed:
            columns.append("seed")
        frame = dataframe[columns].dropna().copy()
        if subject_key is not None:
            # One value per repeated-measures cell prevents trials with more repetitions from
            # receiving disproportionate weight.
            cell_keys = [subject_key, "condition", factor_column]
            seed_cell_keys = [*cell_keys, "seed"] if has_seed else cell_keys
            inconsistent_order = frame.groupby(seed_cell_keys, observed=True)["condition_order"].nunique() > 1
            if inconsistent_order.any():
                raise ValueError(
                    f"condition_order must be constant within participant/condition/{factor_column} cells"
                )
            frame = frame.groupby(seed_cell_keys, as_index=False, observed=True).agg(
                {feature: "mean", "condition_order": "first"}
            )
            if has_seed:
                order_across_seeds = frame.groupby(cell_keys, observed=True)["condition_order"].nunique() > 1
                if order_across_seeds.any():
                    raise ValueError("condition_order must be constant across seeds within participant cells")
                frame = frame.groupby(cell_keys, as_index=False, observed=True).agg(
                    {feature: "mean", "condition_order": "first"}
                )
        if frame["condition"].nunique() < 2 or frame[factor_column].nunique() < 2:
            continue
        design, terms, levels = _factorial_design(frame, subject_key, factor_column)
        response = frame[feature].to_numpy(dtype=float)
        for effect, columns_to_remove in terms.items():
            test = _partial_f_test(design, response, columns_to_remove)
            if test is None:
                raise ValueError(
                    f"factorial effect {effect!r} is not estimable for feature "
                    f"{feature!r}; check participant count, complete {factor_column} cells, "
                    "and counterbalanced condition_order"
                )
            statistic, p_value, numerator_df, denominator_df = test
            rows.append(
                {
                    "feature": feature,
                    "effect": effect,
                    "model": (
                        "participant_fixed_effects_cell_means"
                        if subject_key is not None
                        else "episode_level_factorial_ols"
                    ),
                    "subject_key": subject_key,
                    "n": len(frame),
                    "condition_levels": "/".join(map(str, levels["condition"])),
                    f"{factor_column}_levels": "/".join(map(str, levels[factor_column])),
                    "F": statistic,
                    "df_num": numerator_df,
                    "df_den": denominator_df,
                    "p": p_value,
                }
            )
    return _add_fdr(pd.DataFrame(rows))


def factorial_condition_task(
    dataframe: pd.DataFrame,
    features: list[str],
    *,
    subject_key: str | None = None,
) -> pd.DataFrame:
    """Protocol-v2 condition x task model."""
    return factorial_condition_factor(
        dataframe,
        features,
        subject_key=subject_key,
        factor_column="task_id",
    )


def export_excel(
    path: str,
    comparison: pd.DataFrame,
    success: dict[str, object] | None,
    factorial: pd.DataFrame,
    factor_column: str = "task_id",
) -> None:
    """Write a readable, highlighted multi-sheet workbook."""
    from openpyxl.styles import Font, PatternFill

    significant_fill = PatternFill("solid", fgColor="C6EFCE")
    effect_fill = PatternFill("solid", fgColor="FFEB9C")
    bold = Font(bold=True)

    def autosize(worksheet: object) -> None:
        for column in worksheet.columns:
            width = max(
                (len(str(cell.value)) for cell in column if cell.value is not None),
                default=10,
            )
            worksheet.column_dimensions[column[0].column_letter].width = min(width + 2, 36)

    def highlight(
        worksheet: object,
        significance_column: str,
        effect_column: str | None = None,
    ) -> None:
        headers = {cell.value: cell.column for cell in worksheet[1]}
        significance_index = headers.get(significance_column)
        effect_index = headers.get(effect_column) if effect_column else None
        for row in range(2, worksheet.max_row + 1):
            if significance_index and worksheet.cell(row, significance_index).value in (True, "True"):
                for column in range(1, worksheet.max_column + 1):
                    worksheet.cell(row, column).fill = significant_fill
            if effect_index:
                value = worksheet.cell(row, effect_index).value
                if isinstance(value, (int, float)) and abs(value) >= 0.8:
                    worksheet.cell(row, effect_index).fill = effect_fill
                    worksheet.cell(row, effect_index).font = bold
        autosize(worksheet)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        comparison.to_excel(writer, sheet_name="continuous", index=False)
        if success:
            pd.DataFrame([success]).to_excel(writer, sheet_name="success", index=False)
        factorial_sheet = f"condition_x_{factor_column}"[:31]
        if len(factorial):
            factorial.to_excel(writer, sheet_name=factorial_sheet, index=False)
        if len(comparison):
            highlight(writer.sheets["continuous"], "significant", effect_column="effect_size")
        if "success" in writer.sheets:
            autosize(writer.sheets["success"])
        if len(factorial):
            highlight(writer.sheets[factorial_sheet], "significant")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--features_a",
        required=True,
        help="A_mobile_colocated feature CSV.",
    )
    parser.add_argument(
        "--features_b",
        required=True,
        help="B_desktop_separated feature CSV.",
    )
    parser.add_argument("--out", default="outputs/comparison.csv")
    parser.add_argument("--excel", default=None)
    parser.add_argument(
        "--paired",
        action="store_true",
        help="Deprecated compatibility flag; formal analysis is already paired by default.",
    )
    parser.add_argument(
        "--independent_diagnostic",
        action="store_true",
        help="Explicitly run non-publishable independent-samples diagnostics.",
    )
    parser.add_argument(
        "--pair_keys",
        default=None,
        help="Comma-separated matched-trial keys if pair_id is unavailable.",
    )
    parser.add_argument(
        "--subject_key",
        default=None,
        help="Participant column for repeated-measures factorial blocking (auto-detected).",
    )
    args = parser.parse_args()
    if args.paired and args.independent_diagnostic:
        raise SystemExit("--paired and --independent_diagnostic are mutually exclusive")
    paired = not args.independent_diagnostic

    df_a = pd.read_csv(args.features_a)
    df_b = pd.read_csv(args.features_b)
    try:
        validate_condition_table(df_a, "A_mobile_colocated")
        validate_condition_table(df_b, "B_desktop_separated")
    except ValueError as exc:
        raise SystemExit(f"Invalid condition input: {exc}") from exc

    pair_keys: list[str] = []
    if paired:
        requested_keys = (
            [key.strip() for key in args.pair_keys.split(",") if key.strip()]
            if args.pair_keys
            else ["participant_id"]
        )
        if requested_keys != ["participant_id"]:
            raise SystemExit(
                "Formal paired inference must use --pair_keys participant_id after "
                "within-participant episode aggregation."
            )
        pair_keys = ["participant_id"]
        print(f"Paired analysis keys: {pair_keys}")

    subject_key = args.subject_key or "participant_id"
    if subject_key != "participant_id":
        raise SystemExit("Formal repeated-measures blocking requires subject_key=participant_id")
    excluded = NON_FEATURE | set(pair_keys)
    if subject_key:
        excluded.add(subject_key)
        print(f"Factorial repeated-measures block: {subject_key}")

    features = [
        column
        for column in df_a.columns
        if column not in excluded
        and column in df_b.columns
        and pd.api.types.is_numeric_dtype(df_a[column])
        and pd.api.types.is_numeric_dtype(df_b[column])
    ]
    test_label = "paired t-test" if paired else "Welch t-test (diagnostic only)"
    print(f"Preparing {len(features)} features | episode rows A={len(df_a)}, B={len(df_b)}")

    try:
        participant_a = aggregate_participant_features(df_a, features)
        participant_b = aggregate_participant_features(df_b, features)
        comparison = continuous_comparison(
            participant_a,
            participant_b,
            features,
            paired=paired,
            pair_keys=pair_keys,
        )
    except ValueError as exc:
        raise SystemExit(f"Invalid participant-level design: {exc}") from exc
    print(f"\n=== Continuous features ({test_label}, FDR-corrected) ===")
    with pd.option_context("display.max_columns", None, "display.width", 240):
        print(comparison.round(4).to_string(index=False) if len(comparison) else "  (no testable features)")

    try:
        success = participant_success_comparison(df_a, df_b, paired=paired)
    except ValueError as exc:
        raise SystemExit(f"Invalid success outcomes: {exc}") from exc
    if success:
        print(f"\n=== Success rate ({success['test']}) ===")
        print(
            f"  SR_A={success['SR_A']:.1%} (n={success['n_a']}) "
            f"SR_B={success['SR_B']:.1%} (n={success['n_b']}) "
            f"p={success['p']:.4g}"
        )

    combined = pd.concat([df_a, df_b], ignore_index=True)
    factor_column = "task_id"
    try:
        factorial = factorial_condition_factor(
            combined,
            features,
            subject_key=subject_key,
            factor_column=factor_column,
        )
    except ValueError as exc:
        raise SystemExit(f"Invalid factorial design: {exc}") from exc
    if len(factorial):
        print(f"\n=== Condition x {factor_column} factorial model (FDR-corrected) ===")
        print(
            f"Participant fixed effects use participant-level condition/{factor_column} cell "
            "means and adjust for condition_order; use a mixed-effects model for "
            "confirmatory inference if cells are substantially incomplete."
        )
        with pd.option_context("display.max_columns", None, "display.width", 240):
            print(factorial.round(4).to_string(index=False))

    output = Path(args.out)
    output.parent.mkdir(parents=True, exist_ok=True)
    comparison.to_csv(output, index=False)
    success_path = output.with_name(f"{output.stem}_success{output.suffix}")
    if success:
        pd.DataFrame([success]).to_csv(success_path, index=False)
        print(f"Wrote participant-level success inference to {success_path}")
    if len(factorial):
        factorial_path = output.with_name(f"{output.stem}_factorial{output.suffix}")
        factorial.to_csv(factorial_path, index=False)
        print(f"Wrote factorial results to {factorial_path}")
    print(f"\nWrote comparison to {output}")

    if args.excel:
        excel_path = Path(args.excel)
        excel_path.parent.mkdir(parents=True, exist_ok=True)
        export_excel(str(excel_path), comparison, success, factorial, factor_column)
        print(f"Wrote Excel workbook to {excel_path}")

    metadata_path = output.with_name(f"{output.stem}_metadata.json")
    output_paths = {
        "comparison_csv": output,
        "success_csv": success_path if success else None,
        "factorial_csv": output.with_name(f"{output.stem}_factorial{output.suffix}")
        if len(factorial)
        else None,
        "excel": args.excel,
    }
    write_json(
        metadata_path,
        {
            "artifact_type": "wrist_view_presentation_feature_comparison",
            "arguments": vars(args),
            "input_hashes": fingerprint_inputs(
                {"features_a": args.features_a, "features_b": args.features_b}
            ),
            "output_hashes": {name: fingerprint_path(path) for name, path in output_paths.items()},
            "conditions": {
                "A": "A_mobile_colocated",
                "B": "B_desktop_separated",
            },
            "continuous_analysis_unit": "participant_id",
            "episode_rows": {"A": len(df_a), "B": len(df_b)},
            "participants": {
                "A": sorted(df_a["participant_id"].astype(str).unique().tolist()),
                "B": sorted(df_b["participant_id"].astype(str).unique().tolist()),
            },
            "participant_counts": {
                "A": int(df_a["participant_id"].nunique()),
                "B": int(df_b["participant_id"].nunique()),
            },
            "seed_counts": {
                "A": int(df_a["seed"].nunique()) if "seed" in df_a.columns else 0,
                "B": int(df_b["seed"].nunique()) if "seed" in df_b.columns else 0,
            },
            "seed_aggregation": "participant x seed, then participant"
            if "seed" in df_a.columns or "seed" in df_b.columns
            else "participant",
            "paired": paired,
            "diagnostic_only": not paired,
            "pair_keys": pair_keys,
            "continuous_features": features,
            "multiple_testing": "Benjamini-Hochberg over continuous features",
            "algorithm": {
                "aggregation": "episodes within participant/seed, then seeds within participant",
                "continuous_test": "paired t-test" if paired else "Welch t-test diagnostic",
                "effect_size": "Cohen's dz" if paired else "Cohen's d",
                "fdr_method": "Benjamini-Hochberg",
                "alpha": 0.05,
                "success_test": success["test"] if success else None,
            },
            "factorial": f"participant fixed effects + condition*{factor_column} + condition_order",
            "success_test": success["test"] if success else None,
            "success_limitation": success.get("limitation") if success else None,
            "runtime": runtime_environment(),
            "git": git_state(Path(__file__).parent),
        },
    )
    print(f"Wrote reproducibility metadata to {metadata_path}")


if __name__ == "__main__":
    main()
